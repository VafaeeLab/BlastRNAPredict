import os
import pandas as pd
import numpy as np
import openpyxl
from sklearn.linear_model import LogisticRegressionCV, LogisticRegression
from sklearn.metrics import roc_auc_score, accuracy_score, confusion_matrix
from sklearn.model_selection import StratifiedKFold, cross_val_predict

# ── 1. Load data ──────────────────────────────────────────────────────────────
# Resolve path relative to the repository root (two levels up from this script)
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
wb = openpyxl.load_workbook(
    os.path.join(REPO_ROOT, 'data', 'metadata_260525.xlsx'),
    read_only=True
)
ws = wb['BF cohort']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
data = []
for row in ws.iter_rows(min_row=2):
    data.append({h: row[i].value for i, h in enumerate(headers)})
wb.close()

df = pd.DataFrame(data)
df.columns = df.columns.str.strip()

# ── 2. Encode outcome ─────────────────────────────────────────────────────────
df['outcome'] = (df['hCG_Result'] == 'Positive').astype(int)

# ── 3. Train / validation split ───────────────────────────────────────────────
train = df[df['Cohort'] == 'Training'].dropna(subset=['Age', 'outcome'])
test  = df[df['Cohort'] == 'Validation'].dropna(subset=['Age', 'outcome'])

X_train = train[['Age']].values
y_train = train['outcome'].values
X_test  = test[['Age']].values
y_test  = test['outcome'].values

print(f"Training set  : n={len(y_train)}  (pos={y_train.sum()}, neg={(y_train==0).sum()})")
print(f"Validation set: n={len(y_test)}   (pos={y_test.sum()}, neg={(y_test==0).sum()})")
print(f"Age range (train): {X_train.min():.0f} – {X_train.max():.0f}")
print(f"Age range (test) : {X_test.min():.0f} – {X_test.max():.0f}")

# ── 4. Ridge logistic regression — λ via 10-fold CV (binomial deviance) ──────
Cs = np.logspace(-4, 4, 200)

clf_cv_search = LogisticRegressionCV(
    Cs           = Cs,
    cv           = 10,
    penalty      = 'l2',
    scoring      = 'neg_log_loss',
    solver       = 'lbfgs',
    max_iter     = 5000,
    random_state = 42
)
clf_cv_search.fit(X_train, y_train)

optimal_C      = float(clf_cv_search.C_[0])
optimal_lambda = 1.0 / optimal_C

print(f"\nOptimal λ (1/C) : {optimal_lambda:.6f}   (C = {optimal_C:.6f})")
print(f"Intercept       : {clf_cv_search.intercept_[0]:.4f}")
print(f"  Age           : {clf_cv_search.coef_[0][0]:.4f}")

# ── 5. Internal cross-validated metrics (10-fold, using optimal C) ────────────
cv_splitter = StratifiedKFold(n_splits=10, shuffle=True, random_state=42)
clf_fixed   = LogisticRegression(
    C=optimal_C, penalty='l2', solver='lbfgs', max_iter=5000
)

y_prob_cv = cross_val_predict(
    clf_fixed, X_train, y_train, cv=cv_splitter, method='predict_proba'
)[:, 1]
y_pred_cv = (y_prob_cv >= 0.5).astype(int)

auc_cv  = roc_auc_score(y_train, y_prob_cv)
acc_cv  = accuracy_score(y_train, y_pred_cv)
tn, fp, fn, tp = confusion_matrix(y_train, y_pred_cv).ravel()
sens_cv = tp / (tp + fn)
spec_cv = tn / (tn + fp)

# ── 6. External validation metrics ────────────────────────────────────────────
y_prob_ext = clf_cv_search.predict_proba(X_test)[:, 1]
y_pred_ext = clf_cv_search.predict(X_test)

auc_ext  = roc_auc_score(y_test, y_prob_ext)
acc_ext  = accuracy_score(y_test, y_pred_ext)
tn2, fp2, fn2, tp2 = confusion_matrix(y_test, y_pred_ext).ravel()
sens_ext = tp2 / (tp2 + fn2)
spec_ext = tn2 / (tn2 + fp2)

# ── 7. Report ─────────────────────────────────────────────────────────────────
print("\n" + "="*55)
print("  INTERNAL VALIDATION (10-fold CV, n=65 training)")
print("="*55)
print(f"  AUC         : {auc_cv:.3f}")
print(f"  Accuracy    : {acc_cv:.3f}")
print(f"  Sensitivity : {sens_cv:.3f}")
print(f"  Specificity : {spec_cv:.3f}")

print("\n" + "="*55)
print("  EXTERNAL VALIDATION (n=10 validation set)")
print("="*55)
print(f"  AUC         : {auc_ext:.3f}")
print(f"  Accuracy    : {acc_ext:.3f}")
print(f"  Sensitivity : {sens_ext:.3f}")
print(f"  Specificity : {spec_ext:.3f}")

print(f"\n  Confusion matrix (external):")
print(f"    TP={tp2}, FP={fp2}, FN={fn2}, TN={tn2}")
print(f"    Predicted probabilities: {np.round(y_prob_ext, 3)}")
print(f"    True labels:             {y_test}")
